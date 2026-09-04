#!/usr/bin/env python3
"""
Generate a 14-slide professional PowerPoint presentation for:
Consumer Attention Mapping System — AI-Powered Retail Analytics

Dark navy / blue corporate theme matching the RetailEye AI dashboard.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Constants ──────────────────────────────────────────────────────────────────
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Colors
NAVY = RGBColor(0x0F, 0x17, 0x2A)
DARK_BLUE = RGBColor(0x1E, 0x29, 0x3B)
MID_BLUE = RGBColor(0x33, 0x41, 0x55)
LIGHT_BLUE = RGBColor(0x38, 0xBD, 0xF8)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
GRAY = RGBColor(0x94, 0xA3, 0xB8)
LIGHT_GRAY = RGBColor(0xCB, 0xD5, 0xE1)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
RED = RGBColor(0xEF, 0x44, 0x44)
PURPLE = RGBColor(0xA8, 0x55, 0xF7)
YELLOW = RGBColor(0xEA, 0xB3, 0x08)
GOLD = RGBColor(0xF5, 0x9E, 0x0B)

# Screenshot paths
BRAIN = "/Users/ajeetkumar/.gemini/antigravity/brain/9ed3616b-c5b8-42c2-b56f-64bc40b5eabd"
SCREENSHOTS = {
    "login": f"{BRAIN}/media__1785953349717.png",
    "dashboard_kpi": f"{BRAIN}/media__1786555809712.png",
    "dashboard_sessions": f"{BRAIN}/media__1786555985467.png",
    "dashboard_shelf_attention": f"{BRAIN}/media__1786557793460.png",
    "shelves_planogram": f"{BRAIN}/media__1786553398914.png",
    "shelf_form": f"{BRAIN}/media__1786297202077.png",
    "cameras_page": f"{BRAIN}/media__1786644002352.png",
    "behavior_full": f"{BRAIN}/.user_uploaded/media_1787255395757.png",
    "behavior_journey": f"{BRAIN}/.user_uploaded/media_1787255174754.png",
    "behavior_movement_map": f"{BRAIN}/.user_uploaded/media_1787255018760.png",
    "behavior_journeys_table": f"{BRAIN}/.user_uploaded/media_1787254135323.png",
    "product_leaderboard": f"{BRAIN}/.user_uploaded/media_1787915216557.png",
    "product_ai_engine": f"{BRAIN}/.user_uploaded/media_1787505400310.png",
    "gaze_debug": f"{BRAIN}/gaze_debug_45.jpg",
    "cctv_tracking": f"{BRAIN}/.tempmediaStorage/media_9ed3616b-c5b8-42c2-b56f-64bc40b5eabd_1786860366660.jpg",
    "feature_status": f"{BRAIN}/media__1786554852115.png",
    "login_google": f"{BRAIN}/media__1786686997205.png",
    "dashboard_traffic": f"{BRAIN}/.user_uploaded/media_1787916050571.png",
}

FOOTER_TEXT = "Consumer Attention Mapping System | Infosys Springboard GenAI Internship"

# ── Helper Functions ───────────────────────────────────────────────────────────

def add_dark_bg(slide):
    """Add a solid dark navy background to a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = NAVY

def add_footer(slide, slide_num):
    """Add slide number and footer text."""
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(10), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = f"{FOOTER_TEXT}    |    Slide {slide_num}"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    p.alignment = PP_ALIGN.LEFT

def add_title_bar(slide, title_text, subtitle_text=None, y=Inches(0.3)):
    """Add a styled title bar at top of slide."""
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.7), y, Inches(11.5), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT

    # Accent line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), y + Inches(0.65), Inches(2), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_BLUE
    line.line.fill.background()

    if subtitle_text:
        txBox2 = slide.shapes.add_textbox(Inches(0.7), y + Inches(0.75), Inches(11.5), Inches(0.5))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(14)
        p2.font.color.rgb = GRAY
        p2.alignment = PP_ALIGN.LEFT

def add_text_box(slide, left, top, width, height, text, font_size=12, color=LIGHT_GRAY, bold=False, alignment=PP_ALIGN.LEFT):
    """Add a simple text box."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return tf

def add_bullet_list(slide, left, top, width, height, items, font_size=11, color=LIGHT_GRAY, bullet_color=LIGHT_BLUE):
    """Add a bulleted list."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"▸  {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(4)
    return tf

def add_card(slide, left, top, width, height, title, body_lines, title_color=LIGHT_BLUE, bg_color=DARK_BLUE):
    """Add a rounded card with title and body text."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = MID_BLUE
    card.line.width = Pt(1)

    # Title inside card
    txBox = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), width - Inches(0.3), Inches(0.35))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = title_color

    # Body
    txBox2 = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.4), width - Inches(0.3), height - Inches(0.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    for i, line in enumerate(body_lines):
        if i == 0:
            p2 = tf2.paragraphs[0]
        else:
            p2 = tf2.add_paragraph()
        p2.text = line
        p2.font.size = Pt(9)
        p2.font.color.rgb = LIGHT_GRAY
        p2.space_after = Pt(2)

def add_flow_box(slide, left, top, width, height, text, bg_color=DARK_BLUE, text_color=WHITE, font_size=10):
    """Add a flow diagram box."""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = bg_color
    box.line.color.rgb = MID_BLUE
    box.line.width = Pt(1)
    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = text_color
    p.font.bold = True
    return box

def add_arrow_down(slide, cx, top, length=Inches(0.25)):
    """Add a small downward arrow."""
    arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, cx - Inches(0.08), top, Inches(0.16), length)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = LIGHT_BLUE
    arrow.line.fill.background()

def add_screenshot(slide, img_path, left, top, width, height=None):
    """Add a screenshot image if it exists."""
    if os.path.exists(img_path):
        if height:
            slide.shapes.add_picture(img_path, left, top, width, height)
        else:
            slide.shapes.add_picture(img_path, left, top, width=width)
    else:
        # Placeholder
        add_text_box(slide, left, top, width, Inches(1), f"[Screenshot: {os.path.basename(img_path)}]", 9, GRAY)

def add_notes(slide, text):
    """Add speaker notes."""
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text

# ── Presentation Builder ───────────────────────────────────────────────────────

def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]  # Blank layout

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 1 — TITLE
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)

    # Large gradient-like accent bar at top
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = LIGHT_BLUE
    bar.line.fill.background()

    # Title
    add_text_box(slide, Inches(1), Inches(1.8), Inches(11), Inches(1),
                 "Consumer Attention Mapping System", 42, WHITE, True, PP_ALIGN.CENTER)

    # Subtitle
    add_text_box(slide, Inches(1), Inches(3.0), Inches(11), Inches(0.6),
                 "AI-Powered Retail Analytics & Consumer Behavior Intelligence", 20, LIGHT_BLUE, False, PP_ALIGN.CENTER)

    # Accent line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(3.7), Inches(3), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = CYAN
    line.line.fill.background()

    # Tech tags
    tags = "Computer Vision  •  Deep Learning  •  Behavioral Analytics  •  Retail Intelligence"
    add_text_box(slide, Inches(1), Inches(4.0), Inches(11), Inches(0.5),
                 tags, 13, GRAY, False, PP_ALIGN.CENTER)

    # Presenter info
    add_text_box(slide, Inches(1), Inches(5.2), Inches(11), Inches(0.4),
                 "Ajeet Kumar", 18, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(5.7), Inches(11), Inches(0.4),
                 "B.Tech CSE (AI & ML)  |  GKCIET, Malda", 12, GRAY, False, PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(6.1), Inches(11), Inches(0.4),
                 "Infosys Springboard GenAI Internship", 12, LIGHT_BLUE, False, PP_ALIGN.CENTER)

    add_footer(slide, 1)
    add_notes(slide, """SLIDE 1 — TITLE SLIDE
This is the Consumer Attention Mapping System — an AI-powered retail analytics platform.
The project uses Computer Vision (YOLOv8, MediaPipe), Deep Learning, and Behavioral Analytics to understand how shoppers move through retail stores, which products attract attention, and how store layouts can be optimized.
Developed as part of the Infosys Springboard GenAI Internship program.
Presenter: Ajeet Kumar, B.Tech CSE (AI & ML), GKCIET Malda.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 2 — PROBLEM STATEMENT & OBJECTIVE
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Problem Statement & Objective")

    # Problem side
    add_text_box(slide, Inches(0.7), Inches(1.4), Inches(5.5), Inches(0.4),
                 "⚠  THE RETAIL CHALLENGE", 14, ORANGE, True)
    problem_items = [
        "Retailers lack visibility into how shoppers navigate stores",
        "No data on which shelves and products receive visual attention",
        "Dwell time and engagement patterns remain unmeasured",
        "Product placement decisions rely on intuition, not data",
        "Shopping journeys and behavioral patterns are invisible",
        "Missed opportunities for layout and merchandising optimization"
    ]
    add_bullet_list(slide, Inches(0.7), Inches(1.9), Inches(5.5), Inches(3.5), problem_items, 11)

    # Objective side
    add_text_box(slide, Inches(7), Inches(1.4), Inches(5.5), Inches(0.4),
                 "🎯  OUR OBJECTIVE", 14, GREEN, True)
    obj_items = [
        "Build an AI-powered analytics platform using store cameras",
        "Detect and track shoppers using YOLOv8 + ByteTrack",
        "Estimate gaze direction and attention using MediaPipe",
        "Measure dwell time and shelf engagement duration",
        "Map complete shopping journeys and movement paths",
        "Classify consumer behavior into actionable segments",
        "Score product attractiveness using weighted AI models",
        "Generate shelf optimization & placement recommendations"
    ]
    add_bullet_list(slide, Inches(7), Inches(1.9), Inches(5.5), Inches(4), obj_items, 11)

    # Divider line
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(1.5), Inches(0.02), Inches(4.5))
    divider.fill.solid()
    divider.fill.fore_color.rgb = MID_BLUE
    divider.line.fill.background()

    add_footer(slide, 2)
    add_notes(slide, """SLIDE 2 — PROBLEM STATEMENT & OBJECTIVE
The retail industry faces a fundamental data gap: while e-commerce platforms track every click and scroll, physical retailers have almost zero visibility into in-store shopper behavior.
Retailers don't know which shelves attract the most attention, how long shoppers dwell in different zones, which products get picked up vs. just viewed, or what common shopping journeys look like.
Our objective is to build a complete AI-powered system that uses existing store surveillance cameras combined with Computer Vision (YOLOv8 for detection, ByteTrack for tracking, MediaPipe for gaze estimation) to capture and analyze all of this data.
The system then applies behavioral analytics to classify shoppers, score product attractiveness, and generate actionable optimization recommendations for store managers.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 3 — SYSTEM ARCHITECTURE / WORKFLOW
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "System Architecture & Pipeline")

    # Flow diagram - vertical pipeline
    flow_items = [
        ("Retail Camera Feed", DARK_BLUE),
        ("Video Ingestion & Frame Extraction", DARK_BLUE),
        ("YOLOv8 Person & Object Detection", RGBColor(0x1E, 0x3A, 0x5F)),
        ("ByteTrack Multi-Object Tracking", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Gaze / Head Pose Estimation (MediaPipe)", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Behavior Feature Extraction & Aggregation", RGBColor(0x2D, 0x3A, 0x2E)),
        ("Consumer Behavior Intelligence Engine", RGBColor(0x2D, 0x3A, 0x2E)),
        ("Analytics, Heatmaps & Recommendations", RGBColor(0x3A, 0x2D, 0x1E)),
    ]

    box_w = Inches(3.2)
    box_h = Inches(0.45)
    start_x = Inches(0.8)
    start_y = Inches(1.5)
    spacing = Inches(0.6)

    for i, (text, color) in enumerate(flow_items):
        y = start_y + i * spacing
        add_flow_box(slide, start_x, y, box_w, box_h, text, color, WHITE, 9)
        if i < len(flow_items) - 1:
            add_arrow_down(slide, start_x + box_w / 2, y + box_h, Inches(0.15))

    # Milestone labels on the right
    ms_data = [
        (Inches(1.5), "MILESTONE 1", "Foundation & Camera Integration", LIGHT_BLUE, 0, 1),
        (Inches(3.3), "MILESTONE 2", "Detection, Tracking & Attention", CYAN, 2, 4),
        (Inches(5.1), "MILESTONE 3", "Behavioral Intelligence & Optimization", GREEN, 5, 7),
    ]

    for y_pos, label, desc, color, start_i, end_i in ms_data:
        # Bracket line
        bracket = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.2), y_pos, Inches(0.03), Inches(1.1))
        bracket.fill.solid()
        bracket.fill.fore_color.rgb = color
        bracket.line.fill.background()

        add_text_box(slide, Inches(4.4), y_pos + Inches(0.1), Inches(2.5), Inches(0.3), label, 12, color, True)
        add_text_box(slide, Inches(4.4), y_pos + Inches(0.4), Inches(2.5), Inches(0.5), desc, 9, GRAY)

    # Right side - Data layer
    add_text_box(slide, Inches(7.5), Inches(1.3), Inches(5), Inches(0.4),
                 "DATA & STORAGE LAYER", 13, LIGHT_BLUE, True)

    db_cards = [
        ("PostgreSQL", "Users, Stores, Cameras,\nShelves, Products, Sessions"),
        ("MongoDB (Planned)", "Products, Categories,\nCampaigns, Metadata"),
        ("Redis (Architecture)", "Session Cache,\nReal-time Metrics"),
        ("Object Storage", "Videos, Snapshots,\nHeatmap Images"),
    ]
    for i, (title, body) in enumerate(db_cards):
        col = i % 2
        row = i // 2
        x = Inches(7.5) + col * Inches(2.7)
        y = Inches(1.8) + row * Inches(1.4)
        add_card(slide, x, y, Inches(2.5), Inches(1.2), title, body.split("\n"))

    # Frontend layer
    add_text_box(slide, Inches(7.5), Inches(4.8), Inches(5), Inches(0.4),
                 "FRONTEND & VISUALIZATION", 13, LIGHT_BLUE, True)
    fe_items = ["React.js Dashboard", "Store & Shelf Management UI", "Behavior Analytics Panels",
                "Product Attractiveness Leaderboard", "Interactive Heatmaps & Journey Maps"]
    add_bullet_list(slide, Inches(7.5), Inches(5.3), Inches(5), Inches(2), fe_items, 10)

    add_footer(slide, 3)
    add_notes(slide, """SLIDE 3 — SYSTEM ARCHITECTURE
This slide shows the complete end-to-end pipeline of the Consumer Attention Mapping System.

The pipeline flows from top to bottom:
1. Retail Camera Feed → Video frames are ingested from store surveillance cameras
2. YOLOv8 Detection → Each frame is processed to detect all persons/shoppers
3. ByteTrack Tracking → Detected persons are assigned persistent IDs across frames
4. MediaPipe Gaze → Head pose and gaze direction are estimated for attention analysis
5. Feature Extraction → Raw tracking data is aggregated into behavioral features
6. Behavior Intelligence → Shoppers are segmented and classified
7. Analytics → Heatmaps, scores, and recommendations are generated

The right side shows the data storage layer (PostgreSQL as primary database) and the React.js frontend dashboard.

Milestones 1, 2, and 3 are clearly mapped to different stages of the pipeline.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 4 — TECHNOLOGY STACK
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Technology Stack")

    categories = [
        ("🖥  Backend", ["Python", "FastAPI", "SQLAlchemy", "JWT Auth"], LIGHT_BLUE),
        ("🎨  Frontend", ["React.js", "JavaScript", "Inline CSS (Dark Theme)"], CYAN),
        ("🗄  Database", ["PostgreSQL", "MongoDB (Planned)"], GREEN),
        ("👁  Computer Vision", ["YOLOv8", "OpenCV", "MediaPipe", "ByteTrack"], ORANGE),
        ("📊  Analytics & ML", ["Pandas", "NumPy", "Scikit-learn", "XGBoost"], PURPLE),
        ("📈  Visualization", ["Chart.js", "Plotly", "Matplotlib"], YELLOW),
        ("🛠  DevOps & Tools", ["Git & GitHub", "Docker", "VS Code", "Postman"], GRAY),
    ]

    for i, (cat_name, techs, color) in enumerate(categories):
        col = i % 4
        row = i // 4
        x = Inches(0.5) + col * Inches(3.2)
        y = Inches(1.5) + row * Inches(2.8)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.9), Inches(2.5))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_BLUE
        card.line.color.rgb = MID_BLUE
        card.line.width = Pt(1)

        add_text_box(slide, x + Inches(0.15), y + Inches(0.1), Inches(2.6), Inches(0.35), cat_name, 13, color, True)

        # Tech items
        for j, tech in enumerate(techs):
            chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                          x + Inches(0.15), y + Inches(0.55) + j * Inches(0.38),
                                          Inches(2.5), Inches(0.32))
            chip.fill.solid()
            chip.fill.fore_color.rgb = RGBColor(0x1A, 0x25, 0x38)
            chip.line.color.rgb = MID_BLUE
            chip.line.width = Pt(0.5)
            tf = chip.text_frame
            tf.paragraphs[0].text = tech
            tf.paragraphs[0].font.size = Pt(10)
            tf.paragraphs[0].font.color.rgb = WHITE
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_footer(slide, 4)
    add_notes(slide, """SLIDE 4 — TECHNOLOGY STACK
Backend: Python with FastAPI framework, SQLAlchemy ORM, JWT authentication.
Frontend: React.js with JavaScript, using inline CSS for the dark navy dashboard theme.
Database: PostgreSQL as the primary relational database. MongoDB is planned for document-style product catalogs.
Computer Vision: YOLOv8 for person/object detection, OpenCV for image processing, MediaPipe for face landmark and gaze estimation, ByteTrack for multi-object tracking.
Analytics: Pandas and NumPy for data processing, Scikit-learn and XGBoost for classification.
Visualization: Chart.js for frontend charts, Plotly and Matplotlib for backend visualizations.
DevOps: Git/GitHub for version control, Docker for containerization, VS Code as IDE, Postman for API testing.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 5 — MILESTONE 1: FOUNDATION
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Milestone 1: Foundation & Core Platform",
                  "Project Initialization  •  Authentication  •  Store Management  •  Camera Integration")

    # Left column - features
    ms1_features = [
        "✅  FastAPI backend with PostgreSQL database",
        "✅  JWT authentication & role-based access control",
        "✅  Google OAuth2 integration",
        "✅  Store management (CRUD operations)",
        "✅  Shelf management with zone configuration",
        "✅  Product catalog management",
        "✅  Camera registration & management",
        "✅  Video upload & MJPEG streaming pipeline",
        "✅  React.js frontend with dark theme dashboard",
        "✅  RESTful API with Swagger documentation",
    ]
    add_bullet_list(slide, Inches(0.7), Inches(1.6), Inches(4.8), Inches(4.5), ms1_features, 11, LIGHT_GRAY)

    # Right - screenshots
    add_screenshot(slide, SCREENSHOTS["login"], Inches(6), Inches(1.4), Inches(3.3), Inches(2.2))
    add_text_box(slide, Inches(6), Inches(3.65), Inches(3.3), Inches(0.3),
                 "🔐  JWT Authentication & Login", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_screenshot(slide, SCREENSHOTS["shelves_planogram"], Inches(9.5), Inches(1.4), Inches(3.3), Inches(2.2))
    add_text_box(slide, Inches(9.5), Inches(3.65), Inches(3.3), Inches(0.3),
                 "🏪  Shelf Management & Planogram", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_screenshot(slide, SCREENSHOTS["cameras_page"], Inches(6), Inches(4.2), Inches(3.3), Inches(2.2))
    add_text_box(slide, Inches(6), Inches(6.45), Inches(3.3), Inches(0.3),
                 "📹  Camera Integration & Video Upload", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_screenshot(slide, SCREENSHOTS["dashboard_kpi"], Inches(9.5), Inches(4.2), Inches(3.3), Inches(2.2))
    add_text_box(slide, Inches(9.5), Inches(6.45), Inches(3.3), Inches(0.3),
                 "📊  Real-Time Analytics Dashboard", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_footer(slide, 5)
    add_notes(slide, """SLIDE 5 — MILESTONE 1: FOUNDATION & CORE PLATFORM
Milestone 1 established the complete project foundation:

1. Backend: FastAPI application with SQLAlchemy ORM connected to PostgreSQL database.
2. Authentication: JWT token-based authentication with login/register endpoints. Google OAuth2 callback support. Role-based access control (Admin, Store Manager, Analyst).
3. Store Management: Full CRUD for retail stores with location and metadata.
4. Shelf Management: Shelf configuration linked to stores, with zone coordinates for attention mapping.
5. Product Management: Product catalog with brand, category, shelf assignment.
6. Camera Management: Camera registration, video file upload, and MJPEG streaming endpoint.
7. Frontend: React.js dashboard with dark navy theme, sidebar navigation, and responsive layouts.
8. API Documentation: Auto-generated Swagger/OpenAPI docs at /docs endpoint.

The screenshots show the actual working application: Login page with JWT auth, Shelves page with store planogram, Camera page with video upload, and the analytics dashboard with real-time KPI cards.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 6 — MILESTONE 1: BACKEND ARCHITECTURE
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Backend Architecture & API Design",
                  "FastAPI  •  SQLAlchemy  •  PostgreSQL  •  RESTful Microservice Pattern")

    # Architecture flow
    arch_items = [
        ("React.js Frontend", DARK_BLUE),
        ("Axios HTTP Client", DARK_BLUE),
        ("FastAPI Gateway", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Router Layer", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Service / Business Logic", RGBColor(0x2D, 0x3A, 0x2E)),
        ("SQLAlchemy ORM", RGBColor(0x2D, 0x3A, 0x2E)),
        ("PostgreSQL Database", RGBColor(0x3A, 0x2D, 0x1E)),
    ]
    for i, (text, color) in enumerate(arch_items):
        y = Inches(1.4) + i * Inches(0.7)
        add_flow_box(slide, Inches(0.7), y, Inches(2.8), Inches(0.45), text, color, WHITE, 10)
        if i < len(arch_items) - 1:
            add_arrow_down(slide, Inches(0.7) + Inches(1.4), y + Inches(0.45), Inches(0.2))

    # Router table
    add_text_box(slide, Inches(4.2), Inches(1.3), Inches(4), Inches(0.4),
                 "API ROUTERS", 14, LIGHT_BLUE, True)

    routers = [
        ("auth", "/auth", "Login, Register, JWT Token"),
        ("users", "/users", "User Profile, Settings"),
        ("admin", "/admin", "Admin Dashboard, Management"),
        ("store", "/stores", "Store CRUD Operations"),
        ("shelf", "/shelves", "Shelf CRUD, Zone Config"),
        ("camera", "/cameras", "Camera CRUD, Video Upload"),
        ("product", "/products", "Product Catalog CRUD"),
        ("analytics", "/analytics", "Tracking, Attention, Dwell"),
        ("behavior", "/behavior", "Behavior Intelligence"),
        ("analytics_product", "/analytics/products", "Scoring, Recommendations"),
    ]

    # Table header
    add_text_box(slide, Inches(4.2), Inches(1.8), Inches(1.2), Inches(0.3), "Router", 9, GRAY, True)
    add_text_box(slide, Inches(5.5), Inches(1.8), Inches(1.5), Inches(0.3), "Prefix", 9, GRAY, True)
    add_text_box(slide, Inches(7.0), Inches(1.8), Inches(2), Inches(0.3), "Purpose", 9, GRAY, True)

    for i, (name, prefix, purpose) in enumerate(routers):
        y = Inches(2.15) + i * Inches(0.38)
        row_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.1), y - Inches(0.02), Inches(5), Inches(0.35))
        row_bg.fill.solid()
        row_bg.fill.fore_color.rgb = DARK_BLUE if i % 2 == 0 else NAVY
        row_bg.line.fill.background()
        add_text_box(slide, Inches(4.2), y, Inches(1.2), Inches(0.3), name, 9, LIGHT_BLUE)
        add_text_box(slide, Inches(5.5), y, Inches(1.5), Inches(0.3), prefix, 9, CYAN)
        add_text_box(slide, Inches(7.0), y, Inches(2.2), Inches(0.3), purpose, 9, LIGHT_GRAY)

    # Database models on far right
    add_text_box(slide, Inches(9.8), Inches(1.3), Inches(3), Inches(0.4),
                 "DATABASE MODELS", 14, LIGHT_BLUE, True)

    models = ["User", "Store", "Shelf", "ShelfSnapshot", "Product", "Camera",
              "TrackingSession", "TrackingPoint", "ZoneEvent", "AttentionEvent",
              "DwellEvent", "BehaviorProfile", "ProductInteraction", "ProductScore"]
    for i, m in enumerate(models):
        col = i // 7
        row = i % 7
        x = Inches(9.8) + col * Inches(1.5)
        y = Inches(1.8) + row * Inches(0.42)
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(1.4), Inches(0.35))
        chip.fill.solid()
        chip.fill.fore_color.rgb = DARK_BLUE
        chip.line.color.rgb = MID_BLUE
        chip.line.width = Pt(0.5)
        tf = chip.text_frame
        tf.paragraphs[0].text = m
        tf.paragraphs[0].font.size = Pt(8)
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_footer(slide, 6)
    add_notes(slide, """SLIDE 6 — BACKEND ARCHITECTURE & API DESIGN
The backend follows a layered architecture pattern:
- React Frontend sends HTTP requests via Axios
- FastAPI handles routing, validation, and authentication
- Router layer dispatches to appropriate service functions
- Service layer contains business logic
- SQLAlchemy ORM handles database operations
- PostgreSQL stores all persistent data

10 API routers handle different domains: auth, users, admin, store, shelf, camera, product, analytics, behavior, and product analytics.
14 SQLAlchemy models define the database schema including User, Store, Shelf, Product, Camera, TrackingSession, AttentionEvent, BehaviorProfile, ProductInteraction, ProductScore, etc.
All APIs are documented via FastAPI's auto-generated Swagger UI at /docs.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 7 — CAMERA INGESTION
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Camera Feed Integration & Video Pipeline",
                  "Video Upload  •  MJPEG Streaming  •  Frame Extraction  •  AI Processing")

    # Left - Pipeline
    pipe_items = [
        ("Video File Upload (.mp4)", DARK_BLUE),
        ("FastAPI /cameras/upload Endpoint", DARK_BLUE),
        ("OpenCV VideoCapture", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Frame-by-Frame Extraction", RGBColor(0x1E, 0x3A, 0x5F)),
        ("YOLOv8 + ByteTrack Processing", RGBColor(0x2D, 0x3A, 0x2E)),
        ("JPEG Encoding per Frame", RGBColor(0x2D, 0x3A, 0x2E)),
        ("StreamingResponse (MJPEG)", RGBColor(0x3A, 0x2D, 0x1E)),
        ("React <img> Live Feed Display", RGBColor(0x3A, 0x2D, 0x1E)),
    ]

    for i, (text, color) in enumerate(pipe_items):
        y = Inches(1.5) + i * Inches(0.63)
        add_flow_box(slide, Inches(0.7), y, Inches(3), Inches(0.42), text, color, WHITE, 9)
        if i < len(pipe_items) - 1:
            add_arrow_down(slide, Inches(0.7) + Inches(1.5), y + Inches(0.42), Inches(0.18))

    # Right - screenshot
    add_screenshot(slide, SCREENSHOTS["cameras_page"], Inches(4.5), Inches(1.5), Inches(5), Inches(3.3))
    add_text_box(slide, Inches(4.5), Inches(4.9), Inches(5), Inches(0.3),
                 "📹  Camera Management Page — Video Upload & Pipeline Feed", 10, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    # CCTV frame
    add_screenshot(slide, SCREENSHOTS["cctv_tracking"], Inches(10), Inches(1.5), Inches(2.8), Inches(2))
    add_text_box(slide, Inches(10), Inches(3.55), Inches(2.8), Inches(0.3),
                 "🎯  AI Detection Overlay", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    # Gaze debug
    add_screenshot(slide, SCREENSHOTS["gaze_debug"], Inches(10), Inches(4.1), Inches(2.8), Inches(2))
    add_text_box(slide, Inches(10), Inches(6.15), Inches(2.8), Inches(0.3),
                 "👁  Gaze Estimation Debug", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_footer(slide, 7)
    add_notes(slide, """SLIDE 7 — CAMERA FEED INTEGRATION
The system supports video upload through the Camera management page. Store managers can upload CCTV footage (.mp4 files) which are then processed by the AI pipeline.

Pipeline:
1. Video file is uploaded via the /cameras/upload FastAPI endpoint
2. OpenCV's VideoCapture reads the video frame by frame
3. Each frame passes through YOLOv8 for person detection and ByteTrack for tracking
4. MediaPipe processes detected faces for gaze/head pose estimation
5. Processed frames are JPEG-encoded and served as a multipart MJPEG stream
6. The React frontend displays the live annotated feed

The screenshots show:
- The Camera management page with video upload capability and pipeline feed display
- AI detection overlay showing bounding boxes with product density metrics
- Gaze estimation debug frame showing detected persons with gaze direction vectors

Note: The current implementation uses uploaded video files rather than live IP camera feeds. The architecture supports future integration with RTSP streams from physical retail cameras.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 8 — MILESTONE 2: DETECTION & TRACKING
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Milestone 2: Consumer Detection & Tracking",
                  "YOLOv8  •  ByteTrack  •  MediaPipe  •  Session Management  •  Zone Analytics")

    # Detection pipeline
    detect_items = [
        ("Video Frame Input", DARK_BLUE),
        ("YOLOv8 Person Detection", RGBColor(0x1E, 0x3A, 0x5F)),
        ("ByteTrack ID Assignment", RGBColor(0x1E, 0x3A, 0x5F)),
        ("MediaPipe Gaze Estimation", RGBColor(0x2D, 0x3A, 0x2E)),
        ("Trajectory & Zone Mapping", RGBColor(0x2D, 0x3A, 0x2E)),
        ("Session & Event Storage", RGBColor(0x3A, 0x2D, 0x1E)),
    ]

    for i, (text, color) in enumerate(detect_items):
        y = Inches(1.5) + i * Inches(0.72)
        add_flow_box(slide, Inches(0.7), y, Inches(3), Inches(0.45), text, color, WHITE, 10)
        if i < len(detect_items) - 1:
            add_arrow_down(slide, Inches(0.7) + Inches(1.5), y + Inches(0.45), Inches(0.22))

    # Right side - What's tracked
    add_text_box(slide, Inches(4.5), Inches(1.3), Inches(4), Inches(0.4),
                 "TRACKING DATA GENERATED", 14, LIGHT_BLUE, True)

    tracked_items = [
        ("TrackingSession", "Per-shopper session with entry/exit time"),
        ("TrackingPoint", "X, Y coordinate per frame per shopper"),
        ("ZoneEvent", "Zone entry/exit with dwell duration"),
        ("AttentionEvent", "Shelf gaze duration per session"),
        ("DwellEvent", "Zone-level dwell time aggregation"),
    ]

    for i, (name, desc) in enumerate(tracked_items):
        y = Inches(1.8) + i * Inches(0.7)
        add_card(slide, Inches(4.5), y, Inches(4), Inches(0.6), name, [desc])

    # Dashboard screenshot
    add_screenshot(slide, SCREENSHOTS["dashboard_sessions"], Inches(9), Inches(1.4), Inches(3.8), Inches(2.5))
    add_text_box(slide, Inches(9), Inches(3.95), Inches(3.8), Inches(0.3),
                 "📋  Live Consumer Sessions Table", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_screenshot(slide, SCREENSHOTS["dashboard_shelf_attention"], Inches(9), Inches(4.5), Inches(3.8), Inches(2))
    add_text_box(slide, Inches(9), Inches(6.55), Inches(3.8), Inches(0.3),
                 "📊  Shelf Attention & Zone Traffic", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_footer(slide, 8)
    add_notes(slide, """SLIDE 8 — MILESTONE 2: CONSUMER DETECTION & TRACKING
Milestone 2 implements the core Computer Vision pipeline:

1. YOLOv8 Detection: Each video frame is processed by YOLOv8n to detect all persons. The model outputs bounding boxes with confidence scores.
2. ByteTrack Tracking: Detected persons are assigned persistent unique IDs across frames using ByteTrack's motion-based multi-object tracker.
3. MediaPipe Gaze: For each detected person, MediaPipe FaceMesh extracts 468 face landmarks. Using solvePnP, the system calculates 3D head pose (yaw, pitch, roll) and projects a gaze vector to determine shelf attention.
4. Zone Mapping: Shopper coordinates are mapped to predefined store zones (Entry, Shelf A, Shelf B, Checkout).
5. Session Management: Each tracked shopper generates a TrackingSession with entry/exit timestamps, trajectory points, zone events, and attention events.

The data generated includes:
- TrackingSession: Complete shopper visit record
- TrackingPoint: Per-frame X,Y coordinates for path reconstruction
- ZoneEvent: Zone entry/exit with computed dwell duration
- AttentionEvent: Shelf-specific gaze duration measurements
- DwellEvent: Aggregated zone dwell time analytics""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 9 — SHOPPER JOURNEY ANALYTICS
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Shopper Journey Analytics & Movement Mapping",
                  "Path Reconstruction  •  Route Aggregation  •  Zone Transitions  •  Journey Intelligence")

    # Left - journey screenshot
    add_screenshot(slide, SCREENSHOTS["behavior_journey"], Inches(0.5), Inches(1.4), Inches(6), Inches(4))
    add_text_box(slide, Inches(0.5), Inches(5.5), Inches(6), Inches(0.3),
                 "📊  Customer Journey Analytics Panel — KPIs, Movement Map, Routes & Transitions", 10, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    # Right - explanation
    add_text_box(slide, Inches(7), Inches(1.3), Inches(5.5), Inches(0.4),
                 "AGGREGATE JOURNEY INTELLIGENCE", 13, LIGHT_BLUE, True)

    journey_features = [
        "Total Shoppers tracked across sessions",
        "Average Journey Duration per visit",
        "Average Path Length (zones visited)",
        "Top Customer Routes with frequency %",
        "Zone Transition counts and patterns",
        "Store Movement Map with visit density",
    ]
    add_bullet_list(slide, Inches(7), Inches(1.8), Inches(5.5), Inches(2), journey_features, 11)

    add_text_box(slide, Inches(7), Inches(3.6), Inches(5.5), Inches(0.4),
                 "EXAMPLE ROUTES DISCOVERED", 12, CYAN, True)
    routes = [
        "Entrance → Checkout  (direct purchase — 75%)",
        "Entrance → Shelf A → Checkout  (single browse)",
        "Entrance → Shelf A → Shelf B → Checkout  (exploration)",
        "Entrance → Shelf B → Shelf A → Checkout  (comparison)",
    ]
    add_bullet_list(slide, Inches(7), Inches(4.0), Inches(5.5), Inches(1.5), routes, 10, LIGHT_GRAY)

    add_text_box(slide, Inches(7), Inches(5.5), Inches(5.5), Inches(0.4),
                 "WHY AGGREGATE ANALYTICS?", 12, ORANGE, True)
    why_items = [
        "Individual paths are noisy — aggregation reveals true patterns",
        "Store managers need top routes, not 10,000 individual paths",
        "Enables evidence-based layout and signage decisions",
    ]
    add_bullet_list(slide, Inches(7), Inches(5.9), Inches(5.5), Inches(1), why_items, 10, LIGHT_GRAY)

    add_footer(slide, 9)
    add_notes(slide, """SLIDE 9 — SHOPPER JOURNEY ANALYTICS
This slide showcases the Customer Journey Analytics feature implemented in the Behavior Analytics dashboard.

Key capabilities:
- KPI Summary Cards: Total Shoppers, Average Journey Time, Average Path Length
- Store Movement Map: Visual node diagram showing traffic flow between Entry, Shelf A, Shelf B, and Checkout with visit counts
- Top Customer Routes: Aggregated route patterns with percentage distribution
- Zone Transitions: Count of transitions between each zone pair

The system uses a de-jitter filter (1.0-second dwell threshold) to eliminate noisy bounding-box boundary bouncing that would create false zone transitions.

Aggregate journey analytics is more valuable than displaying individual paths because:
1. Individual tracking paths are noisy and hard to interpret
2. Store managers need actionable patterns, not raw data
3. Route frequency analysis reveals which store layouts work and which create bottlenecks""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 10 — ATTENTION & SHELF ENGAGEMENT
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Attention Analysis & Shelf Engagement",
                  "Gaze Estimation  •  Dwell Time  •  Shelf Attention  •  Traffic Heatmaps")

    # Left - Attention pipeline
    att_items = [
        ("Shopper Detection + Tracking", DARK_BLUE),
        ("Head Pose / Gaze Vector (MediaPipe)", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Shelf / Zone Coordinate Mapping", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Attention Point Calculation", RGBColor(0x2D, 0x3A, 0x2E)),
        ("Attention Duration Aggregation", RGBColor(0x2D, 0x3A, 0x2E)),
        ("Shelf Attention Analytics", RGBColor(0x3A, 0x2D, 0x1E)),
    ]

    for i, (text, color) in enumerate(att_items):
        y = Inches(1.5) + i * Inches(0.72)
        add_flow_box(slide, Inches(0.5), y, Inches(3), Inches(0.45), text, color, WHITE, 9)
        if i < len(att_items) - 1:
            add_arrow_down(slide, Inches(0.5) + Inches(1.5), y + Inches(0.45), Inches(0.22))

    # Attention metrics
    add_text_box(slide, Inches(4), Inches(1.3), Inches(3), Inches(0.4),
                 "ATTENTION METRICS", 13, LIGHT_BLUE, True)
    metrics = [
        "Dwell Time — standing duration per zone",
        "View Duration — direct gaze time on display",
        "Shelf Attention — aggregate gaze per shelf",
        "Product Focus — fixation on individual SKUs",
        "Repeated Attention — return gaze frequency",
    ]
    add_bullet_list(slide, Inches(4), Inches(1.8), Inches(3.2), Inches(2.5), metrics, 10)

    # Screenshots
    add_screenshot(slide, SCREENSHOTS["shelves_planogram"], Inches(4), Inches(4.2), Inches(3.2), Inches(2.2))
    add_text_box(slide, Inches(4), Inches(6.45), Inches(3.2), Inches(0.3),
                 "🗺️  Store Planogram with Attention Heatmap", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_screenshot(slide, SCREENSHOTS["dashboard_shelf_attention"], Inches(7.8), Inches(1.4), Inches(5), Inches(2.8))
    add_text_box(slide, Inches(7.8), Inches(4.25), Inches(5), Inches(0.3),
                 "📊  Shelf Attention Share & Zone Traffic Density Dashboard", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_screenshot(slide, SCREENSHOTS["gaze_debug"], Inches(7.8), Inches(4.8), Inches(2.3), Inches(1.6))
    add_text_box(slide, Inches(7.8), Inches(6.45), Inches(2.3), Inches(0.3),
                 "👁  Gaze Debug Frame", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_screenshot(slide, SCREENSHOTS["cctv_tracking"], Inches(10.3), Inches(4.8), Inches(2.3), Inches(1.6))
    add_text_box(slide, Inches(10.3), Inches(6.45), Inches(2.3), Inches(0.3),
                 "🎯  Shelf Detection HUD", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_footer(slide, 10)
    add_notes(slide, """SLIDE 10 — ATTENTION ANALYSIS & SHELF ENGAGEMENT
This slide explains how shopper attention is captured and analyzed.

Attention Pipeline:
1. Shopper is detected and tracked via YOLOv8 + ByteTrack
2. MediaPipe FaceMesh extracts 468 3D face landmarks
3. solvePnP calculates head pose (yaw, pitch, roll angles)
4. A gaze ray is projected from the head position toward the store space
5. The ray is intersected with defined shelf/zone bounding boxes
6. Attention duration is accumulated per shelf per session

Five key attention metrics are tracked:
- Dwell Time: How long a shopper stands in front of a shelf zone
- View Duration: Direct gaze time focused on a particular display
- Shelf Attention Time: Total accumulated gaze time per shelf
- Product Focus Duration: Fixation time on individual product SKUs
- Repeated Attention: How often a shopper returns their gaze to the same product

The planogram shows a 2D store layout with color-coded attention heatmap overlays.
The dashboard shows shelf attention share (percentage split between shelves) and zone traffic density with visit counts and average dwell times.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 11 — MILESTONE 3: BEHAVIOR INTELLIGENCE
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Milestone 3: Consumer Behavior Intelligence Engine",
                  "Feature Extraction  •  Pattern Analysis  •  Consumer Segmentation  •  Behavior Profiling")

    # Pipeline
    beh_items = [
        ("Raw Milestone 2 Tracking Data", DARK_BLUE),
        ("Behavior Feature Extraction", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Shopping Pattern Analysis", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Consumer Segmentation", RGBColor(0x2D, 0x3A, 0x2E)),
        ("Behavior Profile Generation", RGBColor(0x3A, 0x2D, 0x1E)),
    ]

    for i, (text, color) in enumerate(beh_items):
        y = Inches(1.5) + i * Inches(0.75)
        add_flow_box(slide, Inches(0.5), y, Inches(2.8), Inches(0.45), text, color, WHITE, 9)
        if i < len(beh_items) - 1:
            add_arrow_down(slide, Inches(0.5) + Inches(1.4), y + Inches(0.45), Inches(0.25))

    # Features extracted
    add_text_box(slide, Inches(3.8), Inches(1.3), Inches(3.5), Inches(0.4),
                 "BEHAVIORAL FEATURES EXTRACTED", 12, LIGHT_BLUE, True)

    features = [
        "Visit Duration", "Zones Visited Count", "Products Viewed",
        "Products Picked Up", "Products Returned", "Comparison Count",
        "Attention Duration", "Average Dwell Time", "Movement Distance",
        "Repeat Visit Frequency",
    ]
    for i, feat in enumerate(features):
        col = i % 2
        row = i // 2
        x = Inches(3.8) + col * Inches(1.7)
        y = Inches(1.8) + row * Inches(0.38)
        add_text_box(slide, x, y, Inches(1.6), Inches(0.35), f"▸  {feat}", 9, LIGHT_GRAY)

    # 5 Consumer Segments
    add_text_box(slide, Inches(7.8), Inches(1.3), Inches(5), Inches(0.4),
                 "5 CONSUMER SEGMENTS", 13, GREEN, True)

    segments = [
        ("🔍  Explorer", "High dwell, many zones, broad product exploration", GREEN),
        ("⚡  Quick Buyer", "Short visit, direct to product, fast conversion", CYAN),
        ("⚖️  Comparison Shopper", "Multiple products examined, label reading", PURPLE),
        ("💥  Impulse Buyer", "Sudden gaze shift, promotional display response", ORANGE),
        ("❤️  Brand Loyal", "Repeat brand visits, consistent product preference", RED),
    ]

    for i, (name, desc, color) in enumerate(segments):
        y = Inches(1.8) + i * Inches(1.0)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), y, Inches(5), Inches(0.85))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_BLUE
        card.line.color.rgb = color
        card.line.width = Pt(1.5)
        add_text_box(slide, Inches(7.95), y + Inches(0.08), Inches(4.7), Inches(0.3), name, 12, color, True)
        add_text_box(slide, Inches(7.95), y + Inches(0.4), Inches(4.7), Inches(0.4), desc, 9, GRAY)

    add_footer(slide, 11)
    add_notes(slide, """SLIDE 11 — MILESTONE 3: CONSUMER BEHAVIOR INTELLIGENCE ENGINE
The Behavior Intelligence Engine is the first major component of Milestone 3. Its purpose is to convert raw Milestone 2 tracking events into meaningful behavioral intelligence.

Pipeline:
1. Raw tracking data (positions, zone events, attention events) from Milestone 2
2. Feature Extraction: Compute behavioral features like visit duration, zones visited, products viewed/picked/returned, attention duration, movement distance
3. Pattern Analysis: Analyze shopping patterns, product preferences, and movement behaviors
4. Consumer Segmentation: Classify each shopper into one of 5 behavioral segments
5. Profile Generation: Create a comprehensive BehaviorProfile stored in the database

The 5 consumer segments defined by the project specification:
- Explorer: Broad browsing across many zones and products
- Quick Buyer: Direct, purposeful shopping with fast conversion
- Comparison Shopper: Examines multiple competing products before deciding
- Impulse Buyer: Responds to promotional displays and visual cues
- Brand Loyal: Consistently returns to specific brands across visits

Current implementation uses rule-based segmentation logic. The architecture supports extension to ML-based clustering (K-means, DBSCAN) or classification (XGBoost) as training data accumulates.""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 12 — CONSUMER SEGMENTATION UI
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Consumer Segmentation & Behavior Analytics Dashboard",
                  "Journey Visualization  •  Segment Classification  •  Movement Analysis")

    # Left - full behavior screenshot
    add_screenshot(slide, SCREENSHOTS["behavior_full"], Inches(0.5), Inches(1.4), Inches(6.5), Inches(4.3))
    add_text_box(slide, Inches(0.5), Inches(5.8), Inches(6.5), Inches(0.3),
                 "📊  Full Behavior Analytics Dashboard — Movement Map, Routes, Zone Transitions", 10, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    # Right - journey table
    add_screenshot(slide, SCREENSHOTS["behavior_journeys_table"], Inches(7.5), Inches(1.4), Inches(5.3), Inches(2.5))
    add_text_box(slide, Inches(7.5), Inches(3.95), Inches(5.3), Inches(0.3),
                 "📋  Shopper Journey Table — Session ID, Segment, Route", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    # Segmentation approach
    add_text_box(slide, Inches(7.5), Inches(4.6), Inches(5.3), Inches(0.4),
                 "SEGMENTATION APPROACH", 12, LIGHT_BLUE, True)
    seg_details = [
        "▸  Rule-based initial implementation",
        "▸  Features: dwell time, zones visited, interactions",
        "▸  Tools: Pandas, NumPy for feature computation",
        "▸  Extensible to ML: Scikit-learn, XGBoost",
        "▸  Each session assigned a behavioral segment",
        "▸  Stored as BehaviorProfile in PostgreSQL",
    ]
    tf = add_text_box(slide, Inches(7.5), Inches(5.0), Inches(5.3), Inches(1.8), "", 10, LIGHT_GRAY)
    for i, item in enumerate(seg_details):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(10)
        p.font.color.rgb = LIGHT_GRAY
        p.space_after = Pt(3)

    add_footer(slide, 12)
    add_notes(slide, """SLIDE 12 — CONSUMER SEGMENTATION & BEHAVIOR ANALYTICS DASHBOARD
This slide showcases the implemented Behavior Analytics dashboard with actual screenshots from the working application.

The dashboard features:
- Store Movement Map: Visual node diagram showing shopper flow between store zones with visit counts (Entry: 290 visits, Shelf A: 51 visits, Shelf B: 55 visits, Checkout: 319 visits)
- Top Customer Routes: Aggregated route patterns with percentage distribution
- Zone Transitions: Matrix of transitions between zones
- Recent Shopper Journeys Table: Individual session records showing Session ID, behavioral segment badge (Quick Buyer, Comparison Shopper, etc.), and route sequence

The segmentation approach:
- Currently implemented as rule-based classification using computed behavioral features
- Features include dwell time, number of zones visited, product interactions, and attention duration
- Each session is classified into one of the 5 defined segments
- The architecture is designed to be extensible to ML-based approaches (K-means clustering, XGBoost classification) as more training data accumulates
- All profiles are stored in the PostgreSQL database as BehaviorProfile records""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 13 — HEATMAP + SCORING + RECOMMENDATIONS
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Milestone 3: Heatmaps, Product Scoring & AI Recommendations",
                  "Attention Heatmaps  •  Weighted Attractiveness Model  •  Optimization Engine")

    # Pipeline at top
    pipe_boxes = [
        ("Behavior\nData", DARK_BLUE),
        ("Heatmap\nGeneration", RGBColor(0x1E, 0x3A, 0x5F)),
        ("Product\nScoring", RGBColor(0x2D, 0x3A, 0x2E)),
        ("AI\nRecommendations", RGBColor(0x3A, 0x2D, 0x1E)),
    ]
    for i, (text, color) in enumerate(pipe_boxes):
        x = Inches(0.5) + i * Inches(3.2)
        add_flow_box(slide, x, Inches(1.3), Inches(2.5), Inches(0.6), text, color, WHITE, 10)
        if i < len(pipe_boxes) - 1:
            # Right arrow
            arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(2.6), Inches(1.45), Inches(0.5), Inches(0.3))
            arr.fill.solid()
            arr.fill.fore_color.rgb = LIGHT_BLUE
            arr.line.fill.background()

    # Scoring formula
    add_text_box(slide, Inches(0.5), Inches(2.2), Inches(6), Inches(0.4),
                 "PRODUCT ATTRACTIVENESS SCORING MODEL", 13, GOLD, True)

    weights = [
        ("Attention Duration", "35%", "Total gaze time on product/shelf"),
        ("Product Interaction", "25%", "Touch, pick-up, examine frequency"),
        ("Product Pickup Rate", "20%", "Ratio of pickups to views"),
        ("Purchase Conversion", "15%", "Ratio of purchases to pickups"),
        ("Repeat Engagement", "5%", "Recurring shopper re-visits"),
    ]

    # Table header
    header_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2.65), Inches(6), Inches(0.35))
    header_bg.fill.solid()
    header_bg.fill.fore_color.rgb = MID_BLUE
    header_bg.line.fill.background()
    add_text_box(slide, Inches(0.6), Inches(2.68), Inches(2.5), Inches(0.3), "Metric", 10, WHITE, True)
    add_text_box(slide, Inches(3.2), Inches(2.68), Inches(0.8), Inches(0.3), "Weight", 10, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, Inches(4.1), Inches(2.68), Inches(2.3), Inches(0.3), "Description", 10, WHITE, True)

    for i, (metric, weight, desc) in enumerate(weights):
        y = Inches(3.05) + i * Inches(0.38)
        row_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), y - Inches(0.02), Inches(6), Inches(0.35))
        row_bg.fill.solid()
        row_bg.fill.fore_color.rgb = DARK_BLUE if i % 2 == 0 else NAVY
        row_bg.line.fill.background()
        add_text_box(slide, Inches(0.6), y, Inches(2.5), Inches(0.3), metric, 10, LIGHT_GRAY)
        add_text_box(slide, Inches(3.2), y, Inches(0.8), Inches(0.3), weight, 10, GOLD, True, PP_ALIGN.CENTER)
        add_text_box(slide, Inches(4.1), y, Inches(2.3), Inches(0.3), desc, 9, GRAY)

    # Recommendations
    add_text_box(slide, Inches(0.5), Inches(5.0), Inches(6), Inches(0.4),
                 "AI OPTIMIZATION RECOMMENDATIONS", 12, GREEN, True)
    rec_items = [
        "🔄  Shelf Optimization — Eye-level placement suggestions",
        "📍  Product Placement — Cross-merchandising positions",
        "🏷️  Promotional Suggestions — End-cap & display targeting",
        "📐  Layout Improvement — Traffic flow optimization",
    ]
    add_bullet_list(slide, Inches(0.5), Inches(5.4), Inches(6), Inches(1.5), rec_items, 10)

    # Right - screenshots
    add_screenshot(slide, SCREENSHOTS["product_ai_engine"], Inches(7), Inches(2.2), Inches(5.8), Inches(2.2))
    add_text_box(slide, Inches(7), Inches(4.45), Inches(5.8), Inches(0.3),
                 "🤖  AI Optimization Engine — Promotional Suggestions & Anomaly Detection", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_screenshot(slide, SCREENSHOTS["product_leaderboard"], Inches(7), Inches(4.9), Inches(5.8), Inches(2))
    add_text_box(slide, Inches(7), Inches(6.95), Inches(5.8), Inches(0.3),
                 "⭐  Product Attractiveness Leaderboard with Weighted Metrics", 9, LIGHT_BLUE, True, PP_ALIGN.CENTER)

    add_footer(slide, 13)
    add_notes(slide, """SLIDE 13 — HEATMAPS, PRODUCT SCORING & AI RECOMMENDATIONS
This slide covers the remaining Milestone 3 components: Heatmaps, Product Attractiveness Scoring, and the AI Recommendation Engine.

Product Attractiveness Scoring:
The scoring engine implements the exact weighted formula from the project specification:
- Attention Duration (35%): Total time shoppers spend gazing at the product/shelf
- Product Interaction Frequency (25%): How often shoppers touch/pick-up/examine
- Product Pickup Rate (20%): Ratio of shoppers picking up vs. merely viewing
- Purchase Conversion Rate (15%): Ratio of interactions leading to purchase
- Repeat Engagement Rate (5%): Frequency of recurring shopper revisits

Final Score = Sum(Metric_i × Weight_i)

AI Recommendation Engine:
The system automatically detects anomalies and generates actionable recommendations:
- High Attention + Low Purchase → Promotional Suggestion (product attracts eyes but not sales)
- Low Attention + High Purchase → Hidden Gem / Repositioning (product sells well despite poor visibility)
- Shelf Optimization → Eye-level placement suggestions based on attention metrics
- Layout Improvement → Traffic flow optimization based on movement patterns

The screenshots show:
- AI Optimization Engine cards with specific product recommendations
- Product Attractiveness Leaderboard with ranked products and weighted metric progress bars""")

    # ═══════════════════════════════════════════════════════════════════════════
    # SLIDE 14 — RESULTS & NEXT DIRECTION
    # ═══════════════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(blank_layout)
    add_dark_bg(slide)
    add_title_bar(slide, "Results, Achievements & Next Direction")

    # Implemented column
    add_text_box(slide, Inches(0.5), Inches(1.3), Inches(4), Inches(0.4),
                 "✅  IMPLEMENTED & WORKING", 14, GREEN, True)

    impl_items = [
        "FastAPI backend + React.js frontend",
        "JWT authentication & Google OAuth2",
        "Role-based access control (RBAC)",
        "Store / Shelf / Product / Camera CRUD",
        "Video upload & MJPEG streaming pipeline",
        "YOLOv8 person detection integration",
        "ByteTrack multi-object tracking",
        "MediaPipe gaze estimation pipeline",
        "Tracking sessions with zone events",
        "Attention event recording & aggregation",
        "Store planogram with attention heatmap",
        "Real-time analytics dashboard with KPIs",
        "Consumer behavior feature extraction",
        "Shopper journey analytics & route mapping",
        "Rule-based consumer segmentation",
    ]
    add_bullet_list(slide, Inches(0.5), Inches(1.8), Inches(4), Inches(5), impl_items, 10)

    # Milestone 3 Intelligence
    add_text_box(slide, Inches(4.8), Inches(1.3), Inches(4), Inches(0.4),
                 "🧠  MILESTONE 3 INTELLIGENCE", 14, LIGHT_BLUE, True)

    ms3_items = [
        "Behavior Intelligence Engine",
        "Shopping pattern analysis",
        "5-category consumer segmentation",
        "Product preference analysis",
        "Movement behavior classification",
        "Store & shelf attention heatmaps",
        "Product Attractiveness Scoring Engine",
        "Weighted scoring model (PDF formula)",
        "AI Recommendation Engine",
        "Anomaly detection (High Attn / Low Purchase)",
        "Product Leaderboard dashboard",
        "Virtual Shelf Thermal Heatmap",
    ]
    add_bullet_list(slide, Inches(4.8), Inches(1.8), Inches(4), Inches(5), ms3_items, 10)

    # Next direction
    add_text_box(slide, Inches(9.2), Inches(1.3), Inches(3.8), Inches(0.4),
                 "🔮  NEXT DIRECTION", 14, ORANGE, True)

    next_items = [
        "Connect real retail IP camera feeds",
        "Improve gaze estimation accuracy",
        "Validate behavior segmentation with ML",
        "Multi-camera shopper re-identification",
        "Expand product interaction detection",
        "Executive reporting & PDF exports",
        "Docker containerization & deployment",
        "Cloud deployment (AWS / Azure)",
        "End-to-end production readiness",
    ]
    add_bullet_list(slide, Inches(9.2), Inches(1.8), Inches(3.8), Inches(4), next_items, 10)

    # Bottom bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.8), SLIDE_WIDTH, Inches(0.04))
    bar.fill.solid()
    bar.fill.fore_color.rgb = LIGHT_BLUE
    bar.line.fill.background()

    add_text_box(slide, Inches(1), Inches(6.85), Inches(11), Inches(0.4),
                 "Thank You  —  Consumer Attention Mapping System  —  Ajeet Kumar", 14, WHITE, False, PP_ALIGN.CENTER)

    add_footer(slide, 14)
    add_notes(slide, """SLIDE 14 — RESULTS, ACHIEVEMENTS & NEXT DIRECTION

Summary of what has been achieved through Milestone 3:

IMPLEMENTED & WORKING:
- Complete project foundation with FastAPI backend and React.js frontend
- Authentication system with JWT tokens and Google OAuth2
- Role-based access control for different user types
- Full CRUD management for Stores, Shelves, Products, and Cameras
- Video upload and MJPEG streaming pipeline for camera feeds
- YOLOv8 integration for person detection
- ByteTrack for persistent multi-object tracking
- MediaPipe FaceMesh for gaze estimation
- Complete tracking session management with zone events and attention events
- Store planogram visualization with attention heatmap overlays
- Real-time analytics dashboard with KPI cards
- Consumer behavior feature extraction and journey analytics
- Rule-based consumer segmentation into 5 categories

MILESTONE 3 INTELLIGENCE:
- Behavior Intelligence Engine for pattern analysis and segmentation
- Product Attractiveness Scoring Engine using the weighted model from the PDF
- AI Recommendation Engine detecting anomalies and generating optimization suggestions
- Product Leaderboard and Virtual Shelf Heatmap dashboards

NEXT DIRECTION:
- Integration with real retail IP cameras via RTSP
- Improved gaze estimation accuracy through model fine-tuning
- ML-based behavior segmentation validation
- Multi-camera re-identification for cross-camera tracking
- Docker containerization and cloud deployment""")

    # ─── Save ──────────────────────────────────────────────────────────────────
    output_path = "/Users/ajeetkumar/Desktop/project/Consumer_Attention_Mapping_System_Presentation.pptx"
    prs.save(output_path)
    print(f"✅ Presentation saved to: {output_path}")
    print(f"📊 Total slides: {len(prs.slides)}")
    return output_path


if __name__ == "__main__":
    build_presentation()
